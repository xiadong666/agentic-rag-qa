"""docx 排版工具：统一封面/标题/正文/代码块/提示框/表格/页码。

用 python-docx 直接排版，主要处理两件事：
1. 中文字体：Word 里中文要用 w:eastAsia 指定（只设 ascii 字体会乱）。
2. 段落底纹：代码块/提示框需要给段落加 w:shd。
"""
from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_EA = "宋体"
HEAD_EA = "微软雅黑"
CODE_FONT = "Consolas"
CODE_EA = "微软雅黑"

C_BLUE = RGBColor(0x1F, 0x4E, 0x79)
C_DARK = RGBColor(0x33, 0x33, 0x33)
C_GRAY = RGBColor(0x80, 0x80, 0x80)
C_GREEN = RGBColor(0x1E, 0x7B, 0x34)
C_ORANGE = RGBColor(0xB4, 0x5F, 0x06)
C_RED = RGBColor(0xC0, 0x39, 0x2B)

FILL_CODE = "F2F2F2"
FILL_NOTE = "FFF7E6"   # 注意（黄）
FILL_TIP = "EAF7EA"    # 提示（绿）
FILL_WARN = "FDECEA"   # 坑（红）


def _set_run_fonts(run, *, ascii_font, ea_font, size, bold=False, color=None, italic=False):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    # 东亚字体
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), ea_font)


def _shade(par_or_cell, fill: str):
    """给段落或表格单元格加底纹。"""
    el = par_or_cell._p if hasattr(par_or_cell, "_p") else par_or_cell._tc
    pr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _keep_with_next(par):
    par.paragraph_format.keep_with_next = True


def new_document():
    doc = Document()
    # 页面与正文默认样式
    sec = doc.sections[0]
    sec.top_margin = Cm(2.3)
    sec.bottom_margin = Cm(2.3)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_EA)

    for name, size, color in (
        ("Heading 1", 16, C_BLUE),
        ("Heading 2", 13.5, C_BLUE),
        ("Heading 3", 12, C_DARK),
    ):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_EA)
    return doc


def rich(par, text: str, *, size=11, bold_all=False, color=None, ea=BODY_EA):
    """支持 **加粗** 与 `代码` 的行内标记。"""
    buf = ""
    i = 0
    tokens = []
    mode = None  # None | bold | code
    while i < len(text):
        if text.startswith("**", i):
            tokens.append((mode, buf)); buf = ""
            mode = None if mode == "bold" else "bold"
            i += 2
        elif text[i] == "`":
            tokens.append((mode, buf)); buf = ""
            mode = None if mode == "code" else "code"
            i += 1
        else:
            buf += text[i]
            i += 1
    tokens.append((mode, buf))
    for m, s in tokens:
        if not s:
            continue
        if m == "code":
            _set_run_fonts(par.add_run(s), ascii_font=CODE_FONT, ea_font=CODE_EA,
                           size=size - 0.5, bold=bold_all or False)
        else:
            _set_run_fonts(par.add_run(s), ascii_font="Calibri", ea_font=ea,
                           size=size, bold=bold_all or m == "bold",
                           color=color)
    return par


def para(doc, text="", *, size=11, bold=False, align=None, color=None,
         space_after=6, indent=False, ea=BODY_EA):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.3
    if indent:
        pf.first_line_indent = Pt(22)  # 首行缩进两字符
    if align is not None:
        p.alignment = align
    if text:
        rich(p, text, size=size, bold_all=bold, color=color, ea=ea)
    return p


def h1(doc, text, page_break_before=True):
    if page_break_before:
        doc.add_page_break()
    p = doc.add_heading(text, level=1)
    _keep_with_next(p)
    return p


def h2(doc, text):
    p = doc.add_heading(text, level=2)
    _keep_with_next(p)
    return p


def h3(doc, text):
    p = doc.add_heading(text, level=3)
    _keep_with_next(p)
    return p


def bullet(doc, text, *, level=0, size=11):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    pf = p.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.3
    rich(p, text, size=size)
    return p


def code(doc, text: str, *, size=9.5):
    """等宽代码块：单段落内多行 + 浅灰底纹。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    pf.left_indent = Cm(0.3)
    _shade(p, FILL_CODE)
    lines = text.rstrip("\n").split("\n")
    for idx, line in enumerate(lines):
        run = p.add_run(line if line else " ")
        _set_run_fonts(run, ascii_font=CODE_FONT, ea_font=CODE_EA,
                       size=size, color=C_DARK)
        if idx != len(lines) - 1:
            run.add_break()
    return p


def code_file(doc, rel_path: str, *, start: int = 1, end: int | None = None,
              root: str | None = None):
    """把项目里的真实代码片段贴进文档（保证与代码一致）。"""
    import os
    base = root or os.path.join(os.path.dirname(__file__), "..", "..")
    full = os.path.join(base, rel_path)
    with open(full, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")
    sel = lines[start - 1: end if end else None]
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    _set_run_fonts(p.add_run(f"📄 代码位置：{rel_path}（第 {start}~{start + len(sel) - 1} 行）"),
                   ascii_font="Calibri", ea_font=HEAD_EA, size=9, color=C_GRAY)
    code(doc, "\n".join(sel))


def callout(doc, kind: str, text: str, *, title: str | None = None):
    """提示框：kind ∈ note(注意,黄) / tip(知识点,绿) / warn(坑,红) / idea(亮点,蓝)"""
    fill = {
        "note": (FILL_NOTE, "⚠ 注意", C_ORANGE),
        "tip": (FILL_TIP, "💡 知识点", C_GREEN),
        "warn": (FILL_WARN, "🚫 常见坑", C_RED),
        "idea": (FILL_CODE, "✨ 亮点/面试点", C_BLUE),
    }[kind]
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.3
    _shade(p, fill[0])
    head = title or fill[1]
    _set_run_fonts(p.add_run(head + "  "), ascii_font="Calibri", ea_font=HEAD_EA,
                   size=10.5, bold=True, color=fill[2])
    rich(p, text, size=10.5)
    return p


def table(doc, headers, rows, *, widths=None, header_fill="DCE6F1", font_size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for j, htext in enumerate(headers):
        hdr[j].text = ""
        p = hdr[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_fonts(p.add_run(str(htext)), ascii_font="Calibri", ea_font=HEAD_EA,
                       size=font_size, bold=True)
        _shade(hdr[j], header_fill)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            rich(p, str(val), size=font_size)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def page_break(doc):
    doc.add_page_break()


def add_page_number_footer(doc, text: str):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_fonts(p.add_run(text), ascii_font="Calibri", ea_font=BODY_EA,
                   size=9, color=C_GRAY)
    p.add_run("  ·  第 ")
    _field(p, "PAGE")
    p.add_run(" 页 / 共 ")
    _field(p, "NUMPAGES")
    p.add_run(" 页")


def _field(par, instr: str):
    r1 = par.add_run()
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve")
    instrText.text = f" {instr} "
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    r1._r.append(fldChar1)
    r2 = par.add_run(); r2._r.append(instrText)
    r3 = par.add_run(); r3._r.append(fldChar2)


def add_toc(doc, hint: str):
    p = doc.add_paragraph()
    _set_run_fonts(p.add_run("目录"), ascii_font="Calibri", ea_font=HEAD_EA,
                   size=16, bold=True, color=C_BLUE)
    para(doc, hint, size=9, color=C_GRAY, space_after=10)
    # Word TOC 域（打开后 Ctrl+A → F9 或右键更新域即可生成）
    par = doc.add_paragraph()
    run = par.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    par._p.append(fld)
    doc.add_page_break()


def qa(doc, num, question, points, *, followups=None, tip=None, tag=None, size=10.5):
    """面试题排版：Q 加粗 + 要点 bullets +（可选）追问与加分回答。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(7)
    pf.space_after = Pt(2)
    pf.keep_with_next = True
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.2
    _set_run_fonts(p.add_run(f"Q{num}　"), ascii_font="Calibri", ea_font=HEAD_EA,
                   size=size + 0.5, bold=True, color=C_BLUE)
    _set_run_fonts(p.add_run(question), ascii_font="Calibri", ea_font=HEAD_EA,
                   size=size + 0.5, bold=True, color=C_DARK)
    if tag:
        _set_run_fonts(p.add_run(f"　〔{tag}〕"), ascii_font="Calibri", ea_font=BODY_EA,
                       size=size - 1.5, color=C_GRAY)
    for text in points:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(1)
        bp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        bp.paragraph_format.line_spacing = 1.2
        rich(bp, text, size=size)
    if followups:
        fp = doc.add_paragraph()
        fp.paragraph_format.space_after = Pt(1)
        fp.paragraph_format.left_indent = Cm(0.5)
        _set_run_fonts(fp.add_run("可能追问："), ascii_font="Calibri", ea_font=HEAD_EA,
                       size=size - 0.5, bold=True, color=C_ORANGE)
        rich(fp, "；".join(followups) + "。", size=size - 0.5)
    if tip:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_after = Pt(4)
        tp.paragraph_format.left_indent = Cm(0.5)
        _shade(tp, FILL_TIP)
        _set_run_fonts(tp.add_run("加分回答："), ascii_font="Calibri", ea_font=HEAD_EA,
                       size=size - 0.5, bold=True, color=C_GREEN)
        rich(tp, tip, size=size - 0.5)
