"""单独生成《如何验证 RAG 知识库问答系统是正确的》Word 文档。

用法：
    .venv\\Scripts\\python docs\\gen\\build_verify.py
输出：
    docs/如何验证RAG知识库问答系统是正确的.docx
"""
from __future__ import annotations

import datetime
import os
import sys

GEN_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(GEN_DIR, "..", ".."))
# 导入路径必须在 import 之前就绪：同目录（style + content_p6）+ 项目 scripts/（新增的 Agentic 章节）
sys.path.insert(0, os.path.join(ROOT, "scripts"))
sys.path.insert(0, GEN_DIR)

from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402

import style as S  # noqa: E402
from content_p6 import add as add_verify  # noqa: E402
from content_p7_agentic import add as add_agentic  # noqa: E402  改造后新增：Agent 决策层验收

OUT = os.path.join(ROOT, "docs", "如何验证RAG知识库问答系统是正确的.docx")

TITLE = "如何验证 RAG 知识库问答系统是正确的"
SUB = "分层验收方法 · 可运行验收脚本 · 量化检索指标 · 边界测试清单 · 面试话术 · 附：Agentic RAG 决策层验收"


def cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    S._set_run_fonts(p.add_run(TITLE), ascii_font="Calibri", ea_font="微软雅黑",
                     size=24, bold=True, color=S.C_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    S._set_run_fonts(p.add_run(SUB), ascii_font="Calibri", ea_font="宋体",
                     size=11.5, color=S.C_GRAY)
    doc.add_paragraph()
    meta = [
        ("适用项目", "Agentic RAG 企业级知识库问答与决策系统"
                     "（FastAPI + LangChain 1.2 + FAISS + sentence-transformers + 自研 Agent 编排）"),
        ("配套脚本", "scripts/verify_all.py（朴素链路分层验收 L0~L7）、"
                     "scripts/verify_agentic.py（Agent 决策层验收 A0~A9）、"
                     "scripts/eval_agentic.py（双链路指标评测）"),
        ("实测环境", "Python 3.10.4 / LangChain 1.2.18 / 4 篇文档（含 1 篇 66 分块长文档） / "
                     "33 题双难度标注集（含 10 道库外幻觉探针）"),
        ("实测结果", "verify_all：PASS 21 / FAIL 0；verify_agentic：PASS 22 / FAIL 0；"
                     "双链路对比：应答准确率 100%、幻觉率 0%、生成层准确率提升 4~11 个百分点"),
        ("生成日期", datetime.date.today().strftime("%Y-%m-%d")),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in meta:
        cells = t.add_row().cells
        S._set_run_fonts(cells[0].paragraphs[0].add_run(k),
                         ascii_font="Calibri", ea_font="微软雅黑", size=10.5, bold=True)
        cells[1].paragraphs[0].add_run(v)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    S._set_run_fonts(p.add_run("结论先行：RAG 的“对”要分层证明，而不是靠“问几句感觉还行”"),
                     ascii_font="Calibri", ea_font="宋体", size=10, color=S.C_GRAY)
    doc.add_page_break()


def main() -> None:
    doc = S.new_document()
    cover(doc)
    add_verify(doc)          # 朴素链路：五层验收 + 12 项清单
    add_agentic(doc)         # Agent 决策层：A0~A9 验收 + 12 项增补清单 + 实测数字
    S.add_page_number_footer(doc, "如何验证 RAG 知识库问答系统是正确的")

    out = OUT
    if "--out" in sys.argv:                     # 受限环境可先写到有权限的目录再拷回
        idx = sys.argv.index("--out")
        if idx + 1 < len(sys.argv):
            out = os.path.abspath(sys.argv[idx + 1])
    doc.save(out)
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    print(f"OK -> {out}")
    print(f"headings={len(heads)} paragraphs={len(doc.paragraphs)} "
          f"tables={len(doc.tables)} size={os.path.getsize(out)} bytes")


if __name__ == "__main__":
    main()
