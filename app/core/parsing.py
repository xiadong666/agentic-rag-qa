"""文档解析层。

输入：磁盘上的原始文件；输出：按页（或整篇）切分的文本块，保留 page 信息用于“引用来源”。

知识点：
1. 每种文件格式一个解析函数 → 本质是“策略模式”，方便以后扩展 pptx/html/扫描件OCR。
2. PDF 用 pypdf 按页抽取文本；Word(docx) 本质是 zip+XML，python-docx 帮你解；
   纯文本/Markdown 要处理编码问题（utf-8 失败回退 gbk —— 中文 Windows 常见坑）。
3. 扫描版 PDF 抽不出文字：extract_text 为空 → 主动抛错提示 OCR，而不是悄悄返回空库。
"""
from __future__ import annotations

from dataclasses import dataclass, field

import docx  # python-docx
from pypdf import PdfReader


class ParseError(Exception):
    """解析失败（携带用户可读信息）。"""


# @dataclass 装饰器用于定义数据类，自动生成 __init__、__repr__ 等方法，方便打印和调试
@dataclass
class ParsedDocument:
    """解析结果。

    pages: 列表，元素为 (page_number 或 None, 文本)。PDF 按页记录页码；
           其它格式 page 为 None（正文中称为“无分页来源”）。
    """
    # default_factory=list 表示默认值为可变对象的工厂函数，这里表示默认值为一个空列表
    pages: list[tuple[int | None, str]] = field(default_factory=list)

    # 计算属性，返回所有文本的长度之和，@property 装饰器用于定义计算属性
    @property
    def char_count(self) -> int:
        return sum(len(t) for _, t in self.pages)

    # 计算属性，返回所有文本拼接成的字符串
    @property
    def full_text(self) -> str:
        return "\n".join(t for _, t in self.pages)


def _read_text_file(path: str) -> str:
    """读取 txt/md 等纯文本，自动处理中文编码。

    知识点：Python 读取文本默认 utf-8；Windows 记事本/老文档常是 gbk。
    errors="ignore" 兜底保证不会因个别坏字符崩溃。
    """
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            # 尝试用不同编码解码，成功则返回结果
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _parse_pdf(path: str) -> ParsedDocument:
    reader = PdfReader(path)
    doc = ParsedDocument()
    for i, page in enumerate(reader.pages, start=1):
        # extract_text() 可能返回 None，所以用 or "" 确保 text 始终是字符串
        text = (page.extract_text() or "").strip()
        if text:
            doc.pages.append((i, text))
    if not doc.pages:
        # 常见情况：扫描件 / 图片型 PDF
        raise ParseError(
            f"PDF 未解析出任何文本（共 {len(reader.pages)} 页），"
            "可能是扫描件/图片型 PDF，需要先 OCR（见文档 Step2 说明）。"
        )
    return doc


def _parse_docx(path: str) -> ParsedDocument:
    """docx 解析：正文段落按顺序拼接。

    注意点：表格、页眉页脚、文本框内容不在 document.paragraphs 里，
    真实项目可用 docx 的 block 迭代器进一步处理。
    """
    document = docx.Document(path)
    parts: list[str] = []
    # 遍历文档中的每个段落,document.paragraphs 是一个列表，每个元素是一个段落对象
    for para in document.paragraphs:
        # 获取段落文本并去除首尾空白字符
        text = para.text.strip()
        # 如果段落文本不为空，则添加到 parts 列表中
        if text:
            parts.append(text)
    # 创建 ParsedDocument 对象并添加页面内容
    doc = ParsedDocument()
    doc.pages.append((None, "\n".join(parts)))
    if not parts:
        raise ParseError("docx 未解析出任何文字（文件可能为空文档）。")
    return doc


# 纯文本/Markdown/CSV/JSON/Python/SQL/HTML 解析：直接读取文件内容
def _parse_plain(path: str, ext: str) -> ParsedDocument:
    doc = ParsedDocument()
    doc.pages.append((None, _read_text_file(path)))
    return doc


def parse_file(path: str, filename: str | None = None) -> ParsedDocument:
    """按扩展名分发到具体解析器。

    filename 显式传入可支持“无扩展名但内容明确”的场景（保持健壮）。
    """
    name = (filename or path).lower()
    if name.endswith(".pdf"):
        return _parse_pdf(path)
    if name.endswith(".docx"):
        return _parse_docx(path)
    if name.endswith((".txt", ".md", ".markdown", ".csv", ".json", ".py", ".sql", ".html")):
        return _parse_plain(path, name.rsplit(".", 1)[-1])
    raise ParseError(f"暂不支持的格式：{name.rsplit('.', 1)[-1] if '.' in name else name}")
